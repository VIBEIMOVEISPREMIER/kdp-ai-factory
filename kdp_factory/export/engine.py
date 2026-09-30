from __future__ import annotations
from pathlib import Path
from typing import Iterable
from ..layout.engine import LayoutEngine
class ExportEngine:
    def pdf(self,chapters:list[dict],out:Path,title="",author="",trim_size="8.5x11",bleed=False):
        return LayoutEngine(trim_size,bleed).pdf(chapters,out,title,author)
    def docx(self,chapters:list[dict],out:Path,title="",author=""):
        from docx import Document
        doc=Document()
        if title: doc.add_heading(title,0)
        if author: doc.add_paragraph(author)
        for ch in chapters:
            doc.add_heading(ch.get("title",""),1)
            for p in ch.get("content","").split("\n\n"):
                if p.strip(): doc.add_paragraph(p.strip())
        out.parent.mkdir(parents=True,exist_ok=True); doc.save(out); return out
    def epub(self,chapters:list[dict],out:Path,title="",language="pt-BR"):
        from ebooklib import epub
        book=epub.EpubBook(); book.set_identifier("kdp-ai-factory"); book.set_title(title or "Book"); book.set_language(language)
        items=[]; spine=["nav"]
        for i,ch in enumerate(chapters):
            item=epub.EpubHtml(title=ch.get("title",""),file_name=f"chapter_{i+1}.xhtml",lang=language)
            body="<h1>"+ch.get("title","")+"</h1>"+"".join("<p>"+p.replace("&","&amp;")+"</p>" for p in ch.get("content","").split("\n\n") if p.strip())
            item.content=f"<html><body>{body}</body></html>"; book.add_item(item); items.append(item); spine.append(item)
        book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav()); book.spine=spine
        out.parent.mkdir(parents=True,exist_ok=True); epub.write_epub(str(out),book); return out
