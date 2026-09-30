from __future__ import annotations
from pathlib import Path
from typing import Any
import csv, json, re

class ImportErrorDetail(Exception): pass

def import_file(path: str | Path) -> dict[str, Any]:
    p=Path(path)
    if not p.exists(): raise FileNotFoundError(p)
    ext=p.suffix.lower()
    if ext in {".txt",".md"}: text=p.read_text(encoding="utf-8", errors="replace"); return {"type":"text","source":str(p),"text":text}
    if ext==".csv":
        with p.open("r",encoding="utf-8-sig",newline="") as f: rows=list(csv.DictReader(f))
        return {"type":"table","source":str(p),"rows":rows}
    if ext==".json": return {"type":"json","source":str(p),"data":json.loads(p.read_text(encoding="utf-8"))}
    if ext==".pdf":
        try:
            import fitz
        except ImportError as e: raise ImportErrorDetail("Instale pymupdf para importar PDF.") from e
        doc=fitz.open(p); text="\n".join(page.get_text() for page in doc)
        return {"type":"text","source":str(p),"pages":len(doc),"text":text}
    if ext==".docx":
        try:
            from docx import Document
        except ImportError as e: raise ImportErrorDetail("Instale python-docx para importar DOCX.") from e
        doc=Document(p); text="\n".join(x.text for x in doc.paragraphs)
        return {"type":"text","source":str(p),"text":text}
    if ext==".epub":
        try:
            from ebooklib import epub
            from bs4 import BeautifulSoup
        except ImportError as e: raise ImportErrorDetail("Instale ebooklib e beautifulsoup4 para importar EPUB.") from e
        book=epub.read_epub(str(p)); chunks=[]
        for item in book.get_items():
            if item.get_type()==9: chunks.append(BeautifulSoup(item.get_content(),"html.parser").get_text("\n"))
        return {"type":"text","source":str(p),"text":"\n".join(chunks)}
    if ext==".pptx":
        try:
            from pptx import Presentation
        except ImportError as e: raise ImportErrorDetail("Instale python-pptx para importar PPTX.") from e
        prs=Presentation(p); slides=[]
        for i,s in enumerate(prs.slides,1):
            slides.append({"slide":i,"text":"\n".join(sh.text for sh in s.shapes if hasattr(sh,"text"))})
        return {"type":"slides","source":str(p),"slides":slides}
    if ext in {".png",".jpg",".jpeg",".webp",".bmp"}:
        return {"type":"image","source":str(p),"path":str(p),"mime":f"image/{ext[1:] if ext!='.jpg' else 'jpeg'}"}
    if ext in {".mp3",".wav",".m4a",".ogg"}:
        return {"type":"audio","source":str(p),"path":str(p),"transcription":None}
    raise ImportErrorDetail(f"Formato não suportado: {ext}")

def normalize_to_text(data: dict[str,Any]) -> str:
    if data.get("type")=="text": return data.get("text","")
    if data.get("type")=="table": return "\n".join(" | ".join(map(str,r.values())) for r in data.get("rows",[]))
    if data.get("type")=="slides": return "\n\n".join(s["text"] for s in data.get("slides",[]))
    return ""